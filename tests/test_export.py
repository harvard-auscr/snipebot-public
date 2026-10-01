"""Unit tests for snipebot/export.py: CSV + XLSX (30-aggregate-report.md section 6).

Ledgers are hand-built and run through `eligible_snipes` (never through `parse`),
per the aggregation layer's own testing discipline (see test_aggregate.py).
"""

from __future__ import annotations

import csv

from openpyxl import load_workbook

from snipebot.export import (
    CSV_HEADERS,
    TABLE_ORDER,
    build_all_tables,
    collect_display_ids,
    export_all,
    write_csv,
    write_xlsx,
)

from tests._helpers_export import (
    NAME_CACHE,
    ROS,
    cand,
    elig_of,
    make_resolver,
    mkts,
    selfie,
)


def _sample_elig():
    ledger = [
        cand(mkts(2026, 9, 14, 10, 0, 0), "a1", ("b1",)),
        cand(mkts(2026, 9, 14, 10, 30, 0), "b1", ("a1",)),
        selfie(mkts(2026, 9, 15, 9, 0, 0), "x1", ("a2",)),
        cand(mkts(2026, 9, 14, 9, 0, 0), "a2", ("a1",)),
        # Cooldown reject: same pair inside 15 minutes.
        cand(mkts(2026, 9, 14, 10, 5, 0), "a1", ("b1",)),
    ]
    return elig_of(ledger)


def _build(opted_out=frozenset()):
    elig = _sample_elig()
    tables = build_all_tables(elig, ROS, opted_out)
    resolver = make_resolver()
    names = resolver.resolve_all(collect_display_ids(tables))
    return elig, tables, names


# --------------------------------------------------------------------------- #
# CSV
# --------------------------------------------------------------------------- #

def test_csv_headers_match_spec_order():
    assert CSV_HEADERS["snipes"] == (
        "date", "time", "sniper", "target", "sniper_group", "target_group")
    assert CSV_HEADERS["daily"] == (
        "date", "snipes", "unique_snipers", "unique_targets",
        "cooldown_rejections", "points")
    assert CSV_HEADERS["people"] == (
        "person", "group", "snipes_made", "times_sniped", "unique_targets",
        "unique_snipers", "best_day", "points")
    assert CSV_HEADERS["groups"] == (
        "group", "members", "made", "sniped", "made_per_member",
        "sniped_per_member", "intra_group", "cross_group", "points",
        "points_per_member")
    assert CSV_HEADERS["most_sniped"] == (
        "rank", "person", "group", "times_sniped", "top_sniper_of_them")
    assert CSV_HEADERS["pairs"] == ("sniper", "target", "count", "points")


def test_csv_file_has_bom_and_resolved_names(tmp_path):
    _, tables, names = _build()
    path = tmp_path / "F26_people.csv"
    write_csv(path, "people", tables["people"], names)

    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")

    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    assert rows[0] == list(CSV_HEADERS["people"])
    # Names, not raw Slack IDs, appear in the sniper/target-derived columns.
    people_col = [r[0] for r in rows[1:]]
    assert "Alex" in people_col or "Avery" in people_col
    assert not any(cell.startswith("U") and cell.isupper() for cell in people_col)


def test_csv_best_day_none_is_empty_string(tmp_path):
    _, tables, names = _build()
    path = tmp_path / "F26_people.csv"
    write_csv(path, "people", tables["people"], names)
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    header = rows[0]
    best_day_idx = header.index("best_day")
    # b2 never snipes and is never sniped in this ledger -> excluded entirely
    # (people table only lists people who appear); instead assert any row with
    # snipes_made == 0 has an empty best_day.
    made_idx = header.index("snipes_made")
    for row in rows[1:]:
        if row[made_idx] == "0":
            assert row[best_day_idx] == ""


def test_csv_groups_per_member_dash_when_zero_members(tmp_path):
    _, tables, names = _build()
    path = tmp_path / "F26_groups.csv"
    write_csv(path, "groups", tables["groups"], names)
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    header = rows[0]
    group_idx = header.index("group")
    mpm_idx = header.index("made_per_member")
    for row in rows[1:]:
        if row[group_idx] == "(ungrouped)":
            # ungrouped has no "members" in the roster sense -> per spec, the
            # per-member figure is only a dash when members == 0.
            members_idx = header.index("members")
            if row[members_idx] == "0":
                assert row[mpm_idx] == "—"


def test_csv_bytes_deterministic_across_two_runs(tmp_path):
    _, tables, names = _build()
    p1 = tmp_path / "run1.csv"
    p2 = tmp_path / "run2.csv"
    write_csv(p1, "people", tables["people"], names)
    write_csv(p2, "people", tables["people"], names)
    assert p1.read_bytes() == p2.read_bytes()


def test_csv_points_are_plain_integers(tmp_path):
    _, tables, names = _build()
    path = tmp_path / "F26_daily.csv"
    write_csv(path, "daily", tables["daily"], names)
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    header = rows[0]
    points_idx = header.index("points")
    for row in rows[1:]:
        assert row[points_idx].isdigit()


# --------------------------------------------------------------------------- #
# XLSX
# --------------------------------------------------------------------------- #

def test_xlsx_has_one_sheet_per_table_in_spec_order(tmp_path):
    _, tables, names = _build()
    path = tmp_path / "F26.xlsx"
    write_xlsx(path, tables, names)
    wb = load_workbook(path)
    assert wb.sheetnames == list(TABLE_ORDER)


def test_xlsx_header_row_bold_and_frozen(tmp_path):
    _, tables, names = _build()
    path = tmp_path / "F26.xlsx"
    write_xlsx(path, tables, names)
    wb = load_workbook(path)
    ws = wb["people"]
    assert ws.freeze_panes == "A2"
    for cell in ws[1]:
        assert cell.value is not None
        assert cell.font.bold is True


def test_xlsx_points_column_is_numeric(tmp_path):
    _, tables, names = _build()
    path = tmp_path / "F26.xlsx"
    write_xlsx(path, tables, names)
    wb = load_workbook(path)
    ws = wb["people"]
    header = [c.value for c in ws[1]]
    points_col = header.index("points") + 1
    for row in ws.iter_rows(min_row=2, values_only=False):
        assert isinstance(row[points_col - 1].value, int)


def test_xlsx_points_per_member_is_formatted_string(tmp_path):
    _, tables, names = _build()
    path = tmp_path / "F26.xlsx"
    write_xlsx(path, tables, names)
    wb = load_workbook(path)
    ws = wb["groups"]
    header = [c.value for c in ws[1]]
    ppm_col = header.index("points_per_member") + 1
    for row in ws.iter_rows(min_row=2, values_only=False):
        val = row[ppm_col - 1].value
        assert isinstance(val, str)
        assert val == "—" or "." in val


def test_xlsx_data_row_count_matches_table(tmp_path):
    _, tables, names = _build()
    path = tmp_path / "F26.xlsx"
    write_xlsx(path, tables, names)
    wb = load_workbook(path)
    for table in TABLE_ORDER:
        ws = wb[table]
        assert ws.max_row - 1 == len(tables[table])


# --------------------------------------------------------------------------- #
# export_all
# --------------------------------------------------------------------------- #

# --------------------------------------------------------------------------- #
# Adversarial: every table's CSV, not just the cherry-picked ones above, plus
# a genuine display-name collision surviving the full export.
# --------------------------------------------------------------------------- #

def test_every_table_csv_row_count_matches_table(tmp_path):
    _, tables, names = _build()
    for table in TABLE_ORDER:
        path = tmp_path / f"{table}.csv"
        write_csv(path, table, tables[table], names)
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.reader(fh))
        assert rows[0] == list(CSV_HEADERS[table]), table
        assert len(rows) - 1 == len(tables[table]), table
        for data_row in rows[1:]:
            assert len(data_row) == len(CSV_HEADERS[table]), table


def test_export_all_survives_same_name_same_group_collision(tmp_path):
    from tests._helpers_export import cand, elig_of, mkts
    ledger = [
        cand(mkts(2026, 9, 14, 10, 0, 0), "b1", ("b2",)),
        cand(mkts(2026, 9, 14, 11, 0, 0), "b2", ("b1",)),
    ]
    elig = elig_of(ledger)
    resolver = make_resolver()
    export_all(elig, ROS, frozenset(), resolver, "F26", tmp_path)

    with (tmp_path / "F26_people.csv").open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    person_col = [r[0] for r in rows[1:]]
    # b1 and b2 share both display name and group; the export must still
    # emit two distinct, non-empty identifiers rather than colliding rows.
    assert len(person_col) == len(set(person_col)) == 2


def test_export_all_writes_six_csvs_and_one_xlsx(tmp_path):
    elig = _sample_elig()
    resolver = make_resolver()
    export_all(elig, ROS, frozenset(), resolver, "F26", tmp_path)

    for table in TABLE_ORDER:
        assert (tmp_path / f"F26_{table}.csv").exists()
    assert (tmp_path / "F26.xlsx").exists()


def test_export_all_creates_missing_output_dir(tmp_path):
    elig = _sample_elig()
    resolver = make_resolver()
    out = tmp_path / "nested" / "exports"
    assert not out.exists()
    export_all(elig, ROS, frozenset(), resolver, "F26", out)
    assert out.is_dir()
    assert (out / "F26_snipes.csv").exists()


def test_export_all_deterministic_across_two_runs(tmp_path):
    elig = _sample_elig()
    resolver = make_resolver()
    out1 = tmp_path / "run1"
    out2 = tmp_path / "run2"
    export_all(elig, ROS, frozenset(), resolver, "F26", out1)
    export_all(elig, ROS, frozenset(), resolver, "F26", out2)

    for table in TABLE_ORDER:
        b1 = (out1 / f"F26_{table}.csv").read_bytes()
        b2 = (out2 / f"F26_{table}.csv").read_bytes()
        assert b1 == b2


# --------------------------------------------------------------------------- #
# S19: name cells are text, never a formula/command, and never illegal XML.
# --------------------------------------------------------------------------- #

def test_csv_formula_looking_name_is_prefixed_as_text(tmp_path):
    resolver = make_resolver(cache={**NAME_CACHE, "a1": "=1+1"})
    elig = _sample_elig()
    tables = build_all_tables(elig, ROS, frozenset())
    names = resolver.resolve_all(collect_display_ids(tables))
    path = tmp_path / "F26_people.csv"
    write_csv(path, "people", tables["people"], names)
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    header = rows[0]
    person_idx = header.index("person")
    cells = [r[person_idx] for r in rows[1:]]
    assert "'=1+1" in cells
    assert "=1+1" not in cells


def test_xlsx_formula_looking_name_lands_as_string_cell(tmp_path):
    resolver = make_resolver(cache={**NAME_CACHE, "a1": "=1+1"})
    elig = _sample_elig()
    tables = build_all_tables(elig, ROS, frozenset())
    names = resolver.resolve_all(collect_display_ids(tables))
    path = tmp_path / "F26.xlsx"
    write_xlsx(path, tables, names)
    wb = load_workbook(path)
    ws = wb["people"]
    header = [c.value for c in ws[1]]
    person_col = header.index("person") + 1
    seen = False
    for row in ws.iter_rows(min_row=2, values_only=False):
        cell = row[person_col - 1]
        if cell.value == "=1+1":
            seen = True
            assert cell.data_type == "s"
    assert seen


def test_csv_and_xlsx_name_drops_illegal_control_char(tmp_path):
    resolver = make_resolver(cache={**NAME_CACHE, "a1": "Al\x07ex"})
    elig = _sample_elig()
    tables = build_all_tables(elig, ROS, frozenset())
    names = resolver.resolve_all(collect_display_ids(tables))

    csv_path = tmp_path / "F26_people.csv"
    write_csv(csv_path, "people", tables["people"], names)
    with csv_path.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    person_idx = rows[0].index("person")
    csv_cells_out = [r[person_idx] for r in rows[1:]]
    assert "Alex" in csv_cells_out
    assert not any("\x07" in c for c in csv_cells_out)

    xlsx_path = tmp_path / "F26.xlsx"
    write_xlsx(xlsx_path, tables, names)
    wb = load_workbook(xlsx_path)
    ws = wb["people"]
    header = [c.value for c in ws[1]]
    person_col = header.index("person") + 1
    xlsx_values = [
        row[person_col - 1].value for row in ws.iter_rows(min_row=2, values_only=False)
    ]
    assert "Alex" in xlsx_values
    assert not any("\x07" in v for v in xlsx_values if isinstance(v, str))


def test_csv_name_with_comma_quote_newline_round_trips(tmp_path):
    tricky = 'Al"ex, Jr.\nSecond'
    resolver = make_resolver(cache={**NAME_CACHE, "a1": tricky})
    elig = _sample_elig()
    tables = build_all_tables(elig, ROS, frozenset())
    names = resolver.resolve_all(collect_display_ids(tables))
    path = tmp_path / "F26_people.csv"
    write_csv(path, "people", tables["people"], names)
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    person_idx = rows[0].index("person")
    cells = [r[person_idx] for r in rows[1:]]
    assert tricky in cells


def test_export_all_never_leaks_raw_slack_ids_into_snipes_csv(tmp_path):
    elig = _sample_elig()
    resolver = make_resolver()
    export_all(elig, ROS, frozenset(), resolver, "F26", tmp_path)
    with (tmp_path / "F26_snipes.csv").open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    sniper_idx = rows[0].index("sniper")
    target_idx = rows[0].index("target")
    for row in rows[1:]:
        assert row[sniper_idx] in ("Alex", "Avery", "Ash", "Sam", "Xan") or row[sniper_idx].startswith("[")
        assert row[target_idx] in ("Alex", "Avery", "Ash", "Sam", "Xan") or row[target_idx].startswith("[")
