"""Regression tests for the wave-4 report/export rulings (E-W4-1, 13, 20a, 27, 32).

Each test builds an ``Eligibility`` by hand and renders a digest (or writes an XLSX);
its docstring cites the ruling it pins.
"""

from __future__ import annotations

import csv

from openpyxl import load_workbook

from snipebot.aggregate import Eligibility, EligibleSnipe
from snipebot.config import Section
from snipebot.export import build_all_tables, collect_display_ids, write_csv, write_xlsx
from snipebot.report import (
    MAX_DIGEST_NAME,
    MAX_SECTION_TEXT,
    DigestTooLargeError,
    NameResolver,
    numbers_payload,
    render_digest,
)

from tests._helpers_digest import SEMESTER, TZ, report_of, roster_of

_BASE = 1_789_718_400  # 2026-09-18 08:00:00 UTC
_ALL_RANKED = (Section.DAY, Section.TOP_SNIPERS, Section.MOST_SNIPED, Section.GROUPS,
               Section.PAIRS)


def _snipe(i: int, sniper: str, target: str, groups: dict, *,
           selfie: bool = False) -> EligibleSnipe:
    sec = _BASE + i
    return EligibleSnipe(
        ts=f"{sec}.000000", ts_us=sec * 1_000_000,
        date="2026-09-18", time="08:00:00",
        sniper=sniper, target=target,
        sniper_group=groups.get(sniper), target_group=groups.get(target),
        selfie=selfie,
    )


def _elig(snipes) -> Eligibility:
    return Eligibility(semester=SEMESTER, snipes=tuple(snipes), rejections=())


def _render(report, snipes, groups, cache, *, revision=0):
    roster = roster_of(groups)
    return render_digest(
        report=report, period_key="daily:2026-09-18", semester=SEMESTER,
        elig=_elig(snipes), roster=roster, opted_out=frozenset(),
        names=NameResolver(cache, roster), tz=TZ, revision=revision,
        selfie_emoji="selfie",
    )


def _section_text(render, title: str) -> str:
    for b in render.blocks:
        if b.get("type") == "section" and b["text"]["text"].startswith(title):
            return b["text"]["text"]
    raise AssertionError(f"no section starting {title!r}")


def _mrkdwn_objects(blocks):
    for b in blocks:
        if b["type"] == "section":
            yield b["text"]
            yield from b.get("fields", [])
        elif b["type"] == "context":
            yield from b["elements"]


# --------------------------------------------------------------------------- #
# E-W4-1: every mrkdwn text object is verbatim
# --------------------------------------------------------------------------- #

def test_every_mrkdwn_object_is_verbatim_including_context_and_empty_sections():
    """E-W4-1: every mrkdwn text object a digest emits (section text, fields, the
    `_revised_` and legend context elements, and an empty section's text) sets
    `"verbatim": true`, so no name reaches Slack's mention preprocessing."""
    groups = {"U0AAA001": "sibA", "U0AAA002": "sibB"}
    cache = {"U0AAA001": "user-1 @channel", "U0AAA002": "user-2 @here"}
    snipes = [_snipe(0, "U0AAA001", "U0AAA002", groups, selfie=True)]
    render = _render(report_of(_ALL_RANKED), snipes, groups, cache, revision=1)
    empty = _render(report_of(_ALL_RANKED), [], groups, cache)

    objs = list(_mrkdwn_objects(render.blocks)) + list(_mrkdwn_objects(empty.blocks))
    kinds = {b["type"] for b in render.blocks}
    assert "context" in kinds
    assert objs and all(o["type"] == "mrkdwn" for o in objs)
    assert all(o.get("verbatim") is True for o in objs), [
        o["text"] for o in objs if o.get("verbatim") is not True]


# --------------------------------------------------------------------------- #
# E-W4-13: the XLSX writer drops everything outside the XML 1.0 Char production
# --------------------------------------------------------------------------- #

def test_xlsx_drops_every_non_xml_char_and_the_workbook_opens(tmp_path):
    """E-W4-13: C0 controls other than tab/LF/CR, lone surrogates, U+FFFE and U+FFFF
    are dropped before writing, so one display name never corrupts the workbook."""
    groups = {"U0AAA001": "sibA", "U0AAA002": "sibB", "U0AAA003": "sibB"}
    roster = roster_of(groups)
    elig = _elig([_snipe(0, "U0AAA001", "U0AAA002", groups),
                  _snipe(1, "U0AAA001", "U0AAA003", groups)])
    tables = build_all_tables(elig, roster, frozenset())
    cache = {"U0AAA001": "user-\x011\ud800",
             "U0AAA002": "user-￾2￿",
             "U0AAA003": "user-\udc003\x0b"}
    names = NameResolver(cache, roster).resolve_all(collect_display_ids(tables))
    path = tmp_path / "fall-2026.xlsx"
    write_xlsx(path, tables, names)

    wb = load_workbook(path)
    people = {row[0].value for row in wb["people"].iter_rows(min_row=2)}
    assert {"user-1", "user-2", "user-3"} <= people, people


def test_csv_with_a_lone_surrogate_name_writes_clean_utf8(tmp_path):
    """E-W4-13: the same character filter guards the CSV name cells, so a lone
    surrogate never makes the UTF-8 encoder fail mid-export."""
    groups = {"U0AAA001": "sibA", "U0AAA002": "sibB"}
    roster = roster_of(groups)
    tables = build_all_tables(_elig([_snipe(0, "U0AAA001", "U0AAA002", groups)]),
                              roster, frozenset())
    names = NameResolver({"U0AAA001": "user-1\ud800", "U0AAA002": "user-2"},
                         roster).resolve_all(collect_display_ids(tables))
    path = tmp_path / "people.csv"
    write_csv(path, "people", tables["people"], names)

    with path.open(encoding="utf-8", newline="") as fh:
        cells = [c for row in csv.reader(fh) for c in row]
    assert "user-1" in cells, cells


# --------------------------------------------------------------------------- #
# E-W4-20a: rows <= top_n shows every row; the head never collapses
# --------------------------------------------------------------------------- #

def test_rows_fewer_than_top_n_all_show_with_no_collapse_line():
    """E-W4-20a (30 section 3 case 1): with len(rows) <= top_n every row is head and
    is shown; there is no "...and +K others tied at V" line, even when every row ties."""
    groups = {"U0AAA002": "sibB"}
    cache = {"U0AAA002": "user-2"}
    snipes = []
    for n in range(12):
        uid = f"U0AAA{100 + n}"
        groups[uid] = "sibA"
        cache[uid] = f"user-{100 + n}"
        snipes.append(_snipe(n, uid, "U0AAA002", groups))
    render = _render(report_of((Section.DAY, Section.TOP_SNIPERS), top_n=20),
                     snipes, groups, cache)
    text = _section_text(render, "*Top snipers*")
    assert "others tied at" not in text
    assert all(cache[f"U0AAA{100 + n}"] in text for n in range(12))


def test_head_that_cannot_fit_raises_instead_of_collapsing():
    """E-W4-20a: a head (min(top_n, len(rows)) rows) that cannot fit MAX_SECTION_TEXT
    is a DigestTooLargeError, never a head row hidden under a tie line."""
    groups = {"U0AAA002": "sibB"}
    cache = {"U0AAA002": "user-2"}
    snipes = []
    for n in range(80):
        uid = f"U0AAA{100 + n}"
        groups[uid] = "sibA"
        cache[uid] = (f"user-{100 + n} " + "a" * 40)[:40]
        snipes.append(_snipe(n, uid, "U0AAA002", groups))
    report = report_of((Section.DAY, Section.TOP_SNIPERS), top_n=100)
    try:
        render = _render(report, snipes, groups, cache)
    except DigestTooLargeError:
        return
    raise AssertionError(
        "80 head rows cannot fit 3000 chars, yet the render returned: "
        + _section_text(render, "*Top snipers*")[-80:])


def test_boundary_tie_beyond_top_n_still_collapses():
    """E-W4-20a: rows past top_n tied at V are a boundary tie and still collapse
    into "...and +K others tied at V" when the budget runs out."""
    groups = {"U0AAA002": "sibB"}
    cache = {"U0AAA002": "user-2"}
    snipes = []
    for n in range(80):
        uid = f"U0AAA{100 + n}"
        groups[uid] = "sibA"
        cache[uid] = (f"user-{100 + n} " + "a" * 40)[:40]
        snipes.append(_snipe(n, uid, "U0AAA002", groups))
    render = _render(report_of((Section.DAY, Section.TOP_SNIPERS), top_n=5),
                     snipes, groups, cache)
    text = _section_text(render, "*Top snipers*")
    assert len(text) <= MAX_SECTION_TEXT
    assert text.rsplit("\n", 1)[-1].endswith("others tied at 1"), text[-60:]


# --------------------------------------------------------------------------- #
# E-W4-27: digest names clipped to 40 code points before escaping
# --------------------------------------------------------------------------- #

def test_long_name_is_clipped_to_40_with_ellipsis_inside():
    """E-W4-27: a name longer than 40 code points keeps its first 37 followed by
    `...`; a name of exactly 40 is untouched."""
    long_name = "user-1 " + "b" * 73  # 80 code points, Slack's display-name maximum
    exact = ("user-2 " + "c" * 40)[:40]
    groups = {"U0AAA001": "sibA", "U0AAA002": "sibB"}
    cache = {"U0AAA001": long_name, "U0AAA002": exact}
    render = _render(report_of((Section.DAY, Section.PAIRS)),
                     [_snipe(0, "U0AAA001", "U0AAA002", groups)], groups, cache)
    text = _section_text(render, "*Top pairs*")
    clipped = long_name[:37] + "..."
    assert len(clipped) == MAX_DIGEST_NAME == 40
    assert f"1. {clipped} → {exact} — 1" in text, text
    assert long_name[:38] not in text


def test_clip_counts_code_points_and_escapes_after_clipping():
    """E-W4-27: the clip is by code points (an astral emoji is one) and runs before
    escaping, so `&`/`<` inside the kept part are escaped and an escape entity is
    never cut in half."""
    emoji_name = "\U0001F3AF" * 45  # 45 code points, 90 UTF-16 units
    amp_name = "a" * 35 + "&<x>" + "z" * 10  # the clip keeps "...&<" then "..."
    groups = {"U0AAA001": "sibA", "U0AAA002": "sibB"}
    cache = {"U0AAA001": emoji_name, "U0AAA002": amp_name}
    render = _render(report_of((Section.DAY, Section.PAIRS)),
                     [_snipe(0, "U0AAA001", "U0AAA002", groups)], groups, cache)
    text = _section_text(render, "*Top pairs*")
    assert f"1. {chr(0x1F3AF) * 37}... → " in text, text
    assert "a" * 35 + "&amp;&lt;..." in text, text
    assert "<" not in text and "&lt;x" not in text


def test_disambiguation_suffix_survives_the_clip():
    """E-W4-27: two players sharing a long display name are still told apart by their
    group tag (30 section 4); the base takes the cut, the suffix is kept, and the
    whole rendered name stays within 40 code points."""
    same = "user-dup " + "d" * 60
    groups = {"U0AAA001": "sibA", "U0AAA002": "sibB", "U0AAA003": "sibC"}
    cache = {"U0AAA001": "user-1", "U0AAA002": same, "U0AAA003": same}
    snipes = [_snipe(0, "U0AAA001", "U0AAA002", groups),
              _snipe(1, "U0AAA001", "U0AAA003", groups)]
    render = _render(report_of((Section.DAY, Section.PAIRS)), snipes, groups, cache)
    text = _section_text(render, "*Top pairs*")
    for g in ("sibB", "sibC"):
        name = same[:40 - len(f" ({g})") - 3] + "..." + f" ({g})"
        assert len(name) == 40
        assert f"→ {name} — 1" in text, text


# --------------------------------------------------------------------------- #
# E-W4-32: zero-point rows never appear in a ranked digest table
# --------------------------------------------------------------------------- #

def test_zero_point_rows_never_appear_in_ranked_sections():
    """E-W4-32: a target who made no snipe (0 pts) never appears in top_snipers and a
    sniper never sniped never appears in most_sniped, even when top_n leaves room for
    them. The groups table is exempt: it keeps every group with members (30 section
    5.4), zero-point groups included."""
    groups = {"U0AAA001": "sibA", "U0AAA002": "sibB", "U0AAA003": "sibB",
              "U0AAA004": "sibC"}
    cache = {"U0AAA001": "user-1", "U0AAA002": "user-2", "U0AAA003": "user-3",
             "U0AAA004": "user-4"}
    snipes = [_snipe(0, "U0AAA001", "U0AAA002", groups),
              _snipe(1, "U0AAA001", "U0AAA003", groups)]
    render = _render(report_of(_ALL_RANKED, top_n=10), snipes, groups, cache)

    top = _section_text(render, "*Top snipers*")
    assert top == "*Top snipers*\n1. user-1 — 2 pts (2 snipes)", top
    most = _section_text(render, "*Most sniped*")
    assert "user-1 —" not in most, most
    grp = _section_text(render, "*Groups*")
    assert "sibB" in grp and "sibC" in grp, grp
    assert grp.startswith("*Groups*\n1. sibA"), grp


def test_zero_point_rows_leave_the_numbers_hash_window():
    """E-W4-32: the numbers hash covers the rendered window, so it never carries a
    zero-point row either (a zero row appearing or vanishing cannot revise a digest)."""
    groups = {"U0AAA001": "sibA", "U0AAA002": "sibB"}
    roster = roster_of(groups)
    report = report_of((Section.TOP_SNIPERS, Section.MOST_SNIPED), top_n=10)
    payload = numbers_payload(report, "daily:2026-09-18", SEMESTER,
                              _elig([_snipe(0, "U0AAA001", "U0AAA002", groups)]),
                              roster, frozenset(), TZ)
    assert payload[0] == ["top_snipers", [["U0AAA001", 1, 1]]], payload
    assert payload[1] == ["most_sniped", [["U0AAA002", 1, "U0AAA001"]]], payload


def test_section_of_only_zero_rows_renders_the_empty_form():
    """E-W4-32: when every row of a person ranking is zero, the section renders its
    empty form (30 section 5.4) rather than a list of zeros."""
    groups = {"U0AAA001": None, "U0AAA002": "sibA"}
    cache = {"U0AAA001": "user-1", "U0AAA002": "user-2"}
    render = _render(report_of((Section.DAY, Section.TOP_SNIPERS)), [], groups, cache)
    assert _section_text(render, "*Top snipers*") == "*Top snipers*\n_No snipes yet._"
