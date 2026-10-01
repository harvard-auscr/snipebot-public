"""Wave 4, round 2 probes for the report/export surface (``snipebot.report`` and
``snipebot.export``).

Each test builds an ``Eligibility`` by hand. A test here is expected to FAIL against
the current code: its docstring states the claim and the spec sentence it rests on.
"""

from __future__ import annotations

from openpyxl import load_workbook

from snipebot.aggregate import Eligibility, EligibleSnipe
from snipebot.config import Section
from snipebot.export import build_all_tables, collect_display_ids, write_xlsx
from snipebot.report import DigestTooLargeError, NameResolver, render_digest

from tests._helpers_digest import SEMESTER, TZ, report_of, roster_of

_BASE = 1_789_718_400  # 2026-09-18 08:00:00 UTC


def _snipe(i: int, sniper: str, target: str, groups: dict) -> EligibleSnipe:
    sec = _BASE + i
    return EligibleSnipe(
        ts=f"{sec}.000000", ts_us=sec * 1_000_000,
        date="2026-09-18", time="08:00:00",
        sniper=sniper, target=target,
        sniper_group=groups.get(sniper), target_group=groups.get(target),
        selfie=False,
    )


def _render(report, snipes, roster, cache):
    elig = Eligibility(semester=SEMESTER, snipes=tuple(snipes), rejections=())
    return render_digest(
        report=report, period_key="daily:2026-09-18", semester=SEMESTER,
        elig=elig, roster=roster, opted_out=frozenset(),
        names=NameResolver(cache, roster), tz=TZ, revision=0, selfie_emoji="selfie",
    )


def _section_text(render, title: str) -> str:
    for b in render.blocks:
        if b.get("type") == "section" and b["text"]["text"].startswith(title):
            return b["text"]["text"]
    raise AssertionError(f"no section starting {title!r}")


def _name80(label: str) -> str:
    """An 80-character display name (Slack's display-name maximum)."""
    return (label + " " + "a" * 80)[:80]


# --------------------------------------------------------------------------- #
# 1. Incomplete repair: tied HEAD rows are still collapsed into "+K others tied"
# --------------------------------------------------------------------------- #

def test_tied_head_rows_are_never_collapsed_when_there_is_no_boundary_tie():
    """Claim: the round-1 repair of `_ranked_block` only raises when the character cut
    falls between rows with DIFFERENT values. When the head (ranks 1..top_n) does not
    fit MAX_SECTION_TEXT but the rows at the cut happen to share the cutoff value, it
    still silently drops head rows and appends "...and +K others tied at V". Here
    len(rows) == top_n == 20 (every target sniped once by one sniper, 80-character
    display names), so per 30 section 3 case 1 ("len(rows) <= top_n -> show all; no
    overflow") there is NO boundary tie at all, yet ~4 head rows vanish under a
    "+4 others tied at 1" line. Violates 30 section 3 ("head = rows[0:top_n] (ranks
    1..top_n, always shown)"; "K = number of boundary-tie rows not individually
    shown") and 30 section 5.5 ("Only boundary-tie rows are ever trimmed ... A
    violated assertion is a bug, surfaced as DigestTooLargeError, never a silently
    truncated digest").
    """
    sniper = "U0AAA001"
    groups = {sniper: "sibA"}
    cache = {sniper: _name80("user-1")}
    snipes = []
    for n in range(20):
        target = f"U0AAA{300 + n}"
        groups[target] = "sibB"
        cache[target] = _name80(f"user-{300 + n}")
        snipes.append(_snipe(n, sniper, target, groups))
    report = report_of((Section.DAY, Section.MOST_SNIPED), top_n=20)

    try:
        render = _render(report, snipes, roster_of(groups), cache)
    except DigestTooLargeError:
        return  # the spec-sanctioned outcome when the head cannot fit
    text = _section_text(render, "*Most sniped*")
    assert "others tied at" not in text, (
        "len(rows) == top_n, so there is no boundary tie, yet head rows were "
        f"collapsed: {text.rsplit(chr(10), 1)[-1]!r}"
    )
    # Digest names are clipped to 40 code points (E-W4-27).
    missing = [t for t in (f"U0AAA{300 + n}" for n in range(20))
               if cache[t][:37] + "..." not in text]
    assert not missing, f"{len(missing)} head rows were silently dropped"


# --------------------------------------------------------------------------- #
# 2. count_intra_group: false collapses a scoring groups section to "No snipes yet"
# --------------------------------------------------------------------------- #

def test_groups_section_not_empty_when_only_intra_group_snipes_counted():
    """Claim: with `players.count_intra_group: false`, a semester whose only COUNTED
    snipes are intra-group (every sib selfie is intra-group by construction) renders
    the digest `groups` section as "*Groups*\\n_No snipes yet._", although a member
    of sibA made a COUNTED snipe and sibA holds points (the numbers_hash payload for
    the same digest carries sibA with points 1). `_ranked_block` tests emptiness on
    the mode-dependent raw columns `made`/`sniped`, which count_intra_group:false
    zeroes for intra snipes. Violates 30 section 5.4 ("the empty form applies only
    when no real group made or received a counted snipe ...; otherwise every
    members > 0 group row is shown") and 30 section 2.4 ("count_intra_group governs
    the raw columns ... only ... It never touches points"; the owner wants selfie
    points "counted in the sibfam's standing").
    """
    groups = {"U0AAA001": "sibA", "U0AAA002": "sibA", "U0AAA003": "sibB"}
    roster = roster_of(groups, count_intra_group=False)
    snipes = [_snipe(0, "U0AAA001", "U0AAA002", groups)]
    cache = {u: f"user-{u[-1]}" for u in groups}
    report = report_of((Section.DAY, Section.GROUPS), top_n=5)

    render = _render(report, snipes, roster, cache)
    text = _section_text(render, "*Groups*")
    assert "_No snipes yet._" not in text, (
        "sibA made a COUNTED snipe and has 1 point, yet the groups section reads "
        f"{text!r}"
    )
    assert "1. sibA" in text and "(1 pts, 0/2 made)" in text, text


# --------------------------------------------------------------------------- #
# 3. XLSX: a display name spelled like an Excel error value becomes an error cell
# --------------------------------------------------------------------------- #

def test_xlsx_name_spelled_like_an_error_code_is_a_string_cell(tmp_path):
    """Claim: `write_xlsx` forces a string cell only for names with a formula lead
    character (= + - @ tab CR). openpyxl binds any value equal to an Excel error code
    ("#N/A", "#REF!", "#DIV/0!", "#NAME?", ...) as an ERROR cell (data_type 'e'), so a
    player whose display name is "#N/A" is exported as the #N/A error value rather than
    as their name (a spreadsheet or pandas reader sees an error / missing value, not
    text). Violates 30 section 6 ("Text cells (injection safety): name cells are text
    ... written as a string cell in XLSX").
    """
    groups = {"U0AAA001": "sibA", "U0AAA002": "sibB"}
    roster = roster_of(groups)
    elig = Eligibility(semester=SEMESTER,
                       snipes=(_snipe(0, "U0AAA001", "U0AAA002", groups),),
                       rejections=())
    tables = build_all_tables(elig, roster, frozenset())
    cache = {"U0AAA001": "#N/A", "U0AAA002": "user-2"}
    names = NameResolver(cache, roster).resolve_all(collect_display_ids(tables))
    path = tmp_path / "fall-2026.xlsx"
    write_xlsx(path, tables, names)

    wb = load_workbook(path)
    ws = wb["people"]
    person_cells = [row[0] for row in ws.iter_rows(min_row=2)]
    target = [c for c in person_cells if c.value == "#N/A"]
    assert target, [c.value for c in person_cells]
    assert all(c.data_type == "s" for c in target), (
        f"name cell stored with data_type {[c.data_type for c in target]} "
        "(an error value), not a string"
    )
