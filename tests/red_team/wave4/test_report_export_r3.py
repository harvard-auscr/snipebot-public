"""Wave 4, round 3 probes for the report/export surface (``snipebot.report`` and
``snipebot.export``).

Each test builds an ``Eligibility`` by hand. A test here is expected to FAIL against
the current code: its docstring states the claim and the spec sentence it rests on.
"""

from __future__ import annotations

from openpyxl import load_workbook

from snipebot.aggregate import Eligibility, EligibleSnipe
from snipebot.config import Section
from snipebot.export import build_all_tables, collect_display_ids, write_xlsx
from snipebot.report import (
    MAX_SECTION_TEXT,
    DigestTooLargeError,
    NameResolver,
    render_digest,
)

from tests._helpers_digest import SEMESTER, TZ, report_of, roster_of

_BASE = 1_789_718_400  # 2026-09-18 08:00:00 UTC
_DAY = 86_400


def _snipe(i: int, sniper: str, target: str, groups: dict, *,
           day_offset: int = 0) -> EligibleSnipe:
    sec = _BASE + day_offset * _DAY + i
    return EligibleSnipe(
        ts=f"{sec}.000000", ts_us=sec * 1_000_000,
        date=f"2026-09-{18 + day_offset:02d}", time="08:00:00",
        sniper=sniper, target=target,
        sniper_group=groups.get(sniper), target_group=groups.get(target),
        selfie=False,
    )


def _render(report, snipes, roster, cache, period_key="daily:2026-09-18"):
    elig = Eligibility(semester=SEMESTER, snipes=tuple(snipes), rejections=())
    return render_digest(
        report=report, period_key=period_key, semester=SEMESTER,
        elig=elig, roster=roster, opted_out=frozenset(),
        names=NameResolver(cache, roster), tz=TZ, revision=0, selfie_emoji="selfie",
    )


def _section_text(render, title: str) -> str:
    for b in render.blocks:
        if b.get("type") == "section" and b["text"]["text"].startswith(title):
            return b["text"]["text"]
    raise AssertionError(f"no section starting {title!r}")


# --------------------------------------------------------------------------- #
# 1. A legitimate boundary tie crashes the digest instead of collapsing one more row
# --------------------------------------------------------------------------- #

def test_boundary_tie_overflow_line_fits_by_collapsing_one_more_row():
    """Claim: when the character budget cuts a boundary tie, `_ranked_block` fills
    rows greedily up to MAX_SECTION_TEXT and only then measures the
    "...and +K others tied at V" line. When the filled text sits within ~26 chars of
    3000, it raises DigestTooLargeError ("cannot fit its overflow line") instead of
    collapsing one more boundary-tie row into K. Here 60 snipers with 37-character
    display names (under half Slack's 80-character maximum) are tied at 1 point, top_n
    5: the 5 head rows fit easily, yet the whole digest raises, sync step 9 exits 1 on
    every run, and the day's digest never posts. About a third of all name lengths
    trigger it whenever a tie overflows by characters. Violates 30 section 5.5
    ("When the next line would overflow, stop and append ...and +K others tied at V
    ... If even the title plus the overflow line cannot fit (unreachable ...), the
    render raises DigestTooLargeError") and 30 section 3 step 3 ("show head, then as
    many boundary-tie rows as the message budget allows, and collapse the rest into
    one line").
    """
    target = "U0AAA002"
    groups = {target: "sibB"}
    cache = {target: "user-2"}
    snipes = []
    for n in range(60):
        uid = f"U0AAA{500 + n}"
        groups[uid] = "sibA"
        cache[uid] = (f"user-{500 + n}" + "a" * 80)[:37]
        snipes.append(_snipe(n, uid, target, groups))
    report = report_of((Section.DAY, Section.TOP_SNIPERS), top_n=5)

    try:
        render = _render(report, snipes, roster_of(groups), cache)
    except DigestTooLargeError as exc:
        raise AssertionError(
            "a boundary tie that fits by collapsing one more row raised "
            f"DigestTooLargeError: {exc}") from exc
    text = _section_text(render, "*Top snipers*")
    assert len(text) <= MAX_SECTION_TEXT, len(text)
    assert text.rsplit("\n", 1)[-1].endswith("others tied at 1"), text[-60:]


# --------------------------------------------------------------------------- #
# 3. Incomplete repair: head rows still collapse when there are fewer rows than top_n
# --------------------------------------------------------------------------- #

def test_head_rows_never_collapsed_when_rows_are_fewer_than_top_n():
    """Claim: the round-2 repair of `_ranked_block` raises only when
    `len(ordered) >= top_n`. With fewer rows than top_n, a character-budget cut that
    falls inside a run of equal values is still turned into "...and +K others tied at
    V". Here 17 targets are each sniped once (80-character display names, Slack's
    maximum) with top_n 20: per 30 section 3 case 1 ("len(rows) <= top_n -> show all;
    no overflow") every row is head and there is no boundary tie, yet the last head
    row vanishes under "...and +1 others tied at 1". Violates 30 section 3 ("head =
    rows[0:top_n] (ranks 1..top_n, always shown)"; "K = number of boundary-tie rows
    not individually shown") and 30 section 5.5 ("A violated assertion is a bug,
    surfaced as DigestTooLargeError, never a silently truncated digest").
    """
    sniper = "U0AAA001"
    groups = {sniper: "sibA"}
    cache = {sniper: ("user-1 " + "a" * 80)[:80]}
    snipes = []
    targets = [f"U0AAA{300 + n}" for n in range(17)]
    for n, uid in enumerate(targets):
        groups[uid] = "sibB"
        cache[uid] = (f"user-{300 + n} " + "a" * 80)[:80]
        snipes.append(_snipe(n, sniper, uid, groups))
    report = report_of((Section.DAY, Section.MOST_SNIPED), top_n=20)

    try:
        render = _render(report, snipes, roster_of(groups), cache)
    except DigestTooLargeError:
        return  # the spec-sanctioned outcome when the head cannot fit
    text = _section_text(render, "*Most sniped*")
    assert "others tied at" not in text, (
        "17 rows <= top_n 20, so there is no boundary tie, yet head rows were "
        f"collapsed: {text.rsplit(chr(10), 1)[-1]!r}")
    # Digest names are clipped to 40 code points (E-W4-27).
    missing = [u for u in targets if cache[u][:37] + "..." not in text]
    assert not missing, f"{len(missing)} head rows were silently dropped"


# --------------------------------------------------------------------------- #
# 4. XLSX: an XML non-character in a display name corrupts the whole workbook
# --------------------------------------------------------------------------- #

def test_xlsx_name_with_xml_noncharacter_still_opens(tmp_path):
    """Claim: `_ILLEGAL_XML_CHARS_RE` drops only C0 control characters. U+FFFE and
    U+FFFF are also illegal in XML 1.0 (outside the Char production), and openpyxl
    writes them as-is. A display name carrying one makes the exported <semester>.xlsx
    malformed: openpyxl cannot load it ("not well-formed (invalid token)") and a
    spreadsheet application reports damaged content, so the whole workbook is lost
    to one name. Violates 30 section 6 ("Characters illegal in XML ... are dropped
    before writing").
    """
    groups = {"U0AAA001": "sibA", "U0AAA002": "sibB"}
    roster = roster_of(groups)
    elig = Eligibility(semester=SEMESTER,
                       snipes=(_snipe(0, "U0AAA001", "U0AAA002", groups),),
                       rejections=())
    tables = build_all_tables(elig, roster, frozenset())
    cache = {"U0AAA001": "user-1￿", "U0AAA002": "user-2"}
    names = NameResolver(cache, roster).resolve_all(collect_display_ids(tables))
    path = tmp_path / "fall-2026.xlsx"
    write_xlsx(path, tables, names)

    wb = load_workbook(path)  # raises ParseError on the malformed sheet XML
    people = [row[0].value for row in wb["people"].iter_rows(min_row=2)]
    assert "user-1" in people, people
